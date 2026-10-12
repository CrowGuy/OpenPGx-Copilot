#!/usr/bin/env python3
"""Create the GWAS profiling review workbook from a report payload.

This module is intentionally limited to workbook authoring. Input validation,
profile merging, and review-rule derivation remain in gwas_profile_report.py.

Python 3.10+. Runtime requirement: openpyxl >= 3.1, < 4.
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


COLORS = {
    "navy": "17324D",
    "teal": "147D92",
    "teal_light": "DDF2F4",
    "blue_light": "E8F0F8",
    "text": "203040",
    "muted": "607080",
    "border": "CFD8E3",
    "white": "FFFFFF",
    "edit": "FFF4CC",
    "high": "FADBD8",
    "medium": "FCE8C3",
    "low": "EAF2F8",
    "approved": "DFF0D8",
    "pending": "FFF4CC",
    "risk": "FCE8C3",
    "na": "E8ECEF",
}

BODY_FONT = "Arial"
HEADER_ROW = 4
DATA_START_ROW = 5


class WorkbookBuildError(Exception):
    """Expected workbook-authoring or preview-rendering error."""


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _style_title(ws: Any, last_column: int) -> None:
    last_letter = get_column_letter(max(1, last_column))
    ws.merge_cells(f"A1:{last_letter}1")
    ws.merge_cells(f"A2:{last_letter}2")
    for cell in ws[1]:
        cell.fill = _fill(COLORS["navy"])
    for cell in ws[2]:
        cell.fill = _fill(COLORS["blue_light"])
    ws["A1"].font = Font(name=BODY_FONT, size=16, bold=True, color=COLORS["white"])
    ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"].font = Font(name=BODY_FONT, size=10, italic=True, color=COLORS["muted"])
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 26


def _style_header(ws: Any, row: int, start_column: int, end_column: int) -> None:
    bottom = Side(style="medium", color=COLORS["navy"])
    white_separator = Side(style="thin", color=COLORS["white"])
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row=row, column=column)
        cell.fill = _fill(COLORS["teal"])
        cell.font = Font(name=BODY_FONT, size=10, bold=True, color=COLORS["white"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(
            left=white_separator if column > start_column else Side(style=None),
            bottom=bottom,
        )
    ws.row_dimensions[row].height = 34


def _table_name(sheet_name: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9_]", "", sheet_name)
    if not stem or stem[0].isdigit():
        stem = f"T{stem}"
    return f"{stem}Table"


def _add_table(ws: Any, reference: str, name: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _set_adaptive_widths(ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
    wide_text = {
        "Top Values",
        "Delimiter Candidates",
        "Pair Observations",
        "Review Flags",
        "Failure Examples",
        "Token Count Distribution",
        "Examples",
        "Mismatch Examples",
        "Evidence",
        "Suggested Action",
        "Decision",
        "Notes",
        "Filename",
    }
    identifying = {"Role", "Column", "Column(s)", "Left Column", "Right Column"}
    for index, header in enumerate(headers, start=1):
        observed = len(str(header))
        for row in rows[:100]:
            value = row[index - 1] if index - 1 < len(row) else None
            observed = max(observed, len(str(value or "").split("\n", 1)[0]))
        width = min(max(observed + 2, 10), 24)
        if header in wide_text:
            width = 38 if header == "Suggested Action" else 30
        if header in identifying:
            width = min(max(observed + 2, 18), 30)
        ws.column_dimensions[get_column_letter(index)].width = width


def _apply_conditional_formatting(
    ws: Any, headers: list[str], data_start_row: int, data_end_row: int
) -> None:
    if data_end_row < data_start_row:
        return
    mappings = {
        "Priority": [
            ("High", COLORS["high"], "922B21", True),
            ("Medium", COLORS["medium"], "9A5B00", False),
            ("Low", COLORS["low"], "24527A", False),
        ],
        "Status": [
            ("Pending", COLORS["pending"], COLORS["text"], False),
            ("Approved", COLORS["approved"], COLORS["text"], False),
            ("Accepted Risk", COLORS["risk"], COLORS["text"], False),
            ("Not Applicable", COLORS["na"], COLORS["text"], False),
        ],
    }
    for header, rules in mappings.items():
        if header not in headers:
            continue
        column = headers.index(header) + 1
        letter = get_column_letter(column)
        target = f"{letter}{data_start_row}:{letter}{data_end_row}"
        for value, fill_color, font_color, bold in rules:
            formula = f'EXACT(${letter}{data_start_row},"{value}")'
            ws.conditional_formatting.add(
                target,
                FormulaRule(
                    formula=[formula],
                    fill=_fill(fill_color),
                    font=Font(name=BODY_FONT, color=font_color, bold=bold),
                ),
            )


def _populate_data_sheet(ws: Any, spec: Mapping[str, Any]) -> None:
    headers = [str(item) for item in spec.get("headers", [])]
    rows = [list(row) for row in spec.get("rows", [])]
    if not headers:
        raise WorkbookBuildError(f"Sheet {ws.title!r} has no headers")

    last_column = len(headers)
    data_end_row = DATA_START_ROW + len(rows) - 1
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLORS["teal"]
    ws["A1"] = spec.get("name", ws.title)
    ws["A2"] = spec.get("description", "")
    _style_title(ws, last_column)

    for column, header in enumerate(headers, start=1):
        ws.cell(row=HEADER_ROW, column=column, value=header)
    _style_header(ws, HEADER_ROW, 1, last_column)

    if rows:
        for row_number, row in enumerate(rows, start=DATA_START_ROW):
            for column in range(1, last_column + 1):
                value = row[column - 1] if column - 1 < len(row) else None
                cell = ws.cell(row=row_number, column=column, value=value)
                cell.font = Font(name=BODY_FONT, size=10, color=COLORS["text"])
                cell.alignment = Alignment(vertical="center")
            ws.row_dimensions[row_number].height = 20
        last_letter = get_column_letter(last_column)
        _add_table(
            ws,
            f"A{HEADER_ROW}:{last_letter}{data_end_row}",
            _table_name(ws.title),
        )
    else:
        cell = ws["A5"]
        cell.value = "No observations for this snapshot."
        cell.font = Font(name=BODY_FONT, size=10, italic=True, color=COLORS["muted"])

    header_index = {header: index + 1 for index, header in enumerate(headers)}
    format_groups = [
        (spec.get("percent_columns", []), "0.00%"),
        (spec.get("integer_columns", []), "#,##0"),
        (spec.get("decimal_columns", []), "0.00"),
    ]
    if rows:
        for columns, number_format in format_groups:
            for header in columns:
                column = header_index.get(header)
                if column is None:
                    continue
                for row_number in range(DATA_START_ROW, data_end_row + 1):
                    ws.cell(row=row_number, column=column).number_format = number_format

        for header in spec.get("wrap_columns", []):
            column = header_index.get(header)
            if column is None:
                continue
            for row_number in range(DATA_START_ROW, data_end_row + 1):
                cell = ws.cell(row=row_number, column=column)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.row_dimensions[row_number].height = max(ws.row_dimensions[row_number].height or 20, 42)

        editable_headers = {
            "Approved Type",
            "Missing Reason",
            "Approved Delimiter",
            "Decision",
            "Reviewer",
            "Notes",
        }
        for header in editable_headers:
            column = header_index.get(header)
            if column is None:
                continue
            for row_number in range(DATA_START_ROW, data_end_row + 1):
                ws.cell(row=row_number, column=column).fill = _fill(COLORS["edit"])

        for header, values in spec.get("validations", {}).items():
            column = header_index.get(header)
            if column is None or not values:
                continue
            escaped = [str(value).replace('"', '""') for value in values]
            formula = '"' + ",".join(escaped) + '"'
            if len(formula) > 255:
                raise WorkbookBuildError(
                    f"Validation list for {ws.title!r}/{header!r} exceeds Excel's 255-character limit"
                )
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            validation.errorTitle = "Invalid value"
            validation.error = "Select a value from the list."
            validation.promptTitle = header
            validation.prompt = "Choose an approved review value."
            validation.showErrorMessage = True
            validation.showInputMessage = True
            ws.add_data_validation(validation)
            target = f"{get_column_letter(column)}{DATA_START_ROW}:{get_column_letter(column)}{data_end_row}"
            validation.add(target)
            for row_number in range(DATA_START_ROW, data_end_row + 1):
                ws.cell(row=row_number, column=column).fill = _fill(COLORS["edit"])

    _apply_conditional_formatting(ws, headers, DATA_START_ROW, data_end_row)
    _set_adaptive_widths(ws, headers, rows)
    ws.freeze_panes = "C5" if last_column >= 2 else "A5"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(last_column)}{max(HEADER_ROW, data_end_row)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{HEADER_ROW}"
    ws.sheet_properties.outlinePr.summaryBelow = True


def _populate_overview(ws: Any, payload: Mapping[str, Any]) -> None:
    overview = payload.get("overview", {})
    metadata = list(overview.get("metadata", []))
    kpis = list(overview.get("kpis", []))
    notes = list(overview.get("notes", []))
    files = list(overview.get("files", []))

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = COLORS["navy"]
    ws["A1"] = payload.get("title", "GWAS Profiling Review")
    ws["A2"] = payload.get("subtitle", "")
    _style_title(ws, 8)

    ws["A4"], ws["B4"] = "Provenance", "Value"
    _style_header(ws, 4, 1, 2)
    for row_number, item in enumerate(metadata, start=5):
        key = item[0] if len(item) > 0 else None
        value = item[1] if len(item) > 1 else None
        ws.cell(row=row_number, column=1, value=key)
        ws.cell(row=row_number, column=2, value=value)
        ws.cell(row=row_number, column=1).fill = _fill(COLORS["blue_light"])
        ws.cell(row=row_number, column=1).font = Font(name=BODY_FONT, bold=True, color=COLORS["text"])
        ws.cell(row=row_number, column=2).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row_number].height = 28

    ws["D4"], ws["E4"] = "KPI", "Value"
    _style_header(ws, 4, 4, 5)
    for row_number, item in enumerate(kpis, start=5):
        label = item[0] if len(item) > 0 else None
        value = item[1] if len(item) > 1 else None
        ws.cell(row=row_number, column=4, value=label)
        ws.cell(row=row_number, column=5, value=value)
        ws.cell(row=row_number, column=4).fill = _fill(COLORS["teal_light"])
        ws.cell(row=row_number, column=4).font = Font(name=BODY_FONT, bold=True, color=COLORS["text"])
        ws.cell(row=row_number, column=5).number_format = "#,##0"

    review_spec = next(
        (item for item in payload.get("sheets", []) if item.get("name") == "Review Queue"),
        None,
    )
    if review_spec and review_spec.get("rows") and "Status" in review_spec.get("headers", []):
        status_column = review_spec["headers"].index("Status") + 1
        status_letter = get_column_letter(status_column)
        last_review_row = HEADER_ROW + len(review_spec["rows"])
        ws["G4"], ws["H4"] = "Review Status", "Count"
        _style_header(ws, 4, 7, 8)
        statuses = ["Pending", "In Review", "Approved", "Accepted Risk", "Not Applicable"]
        for row_number, status in enumerate(statuses, start=5):
            ws.cell(row=row_number, column=7, value=status)
            ws.cell(
                row=row_number,
                column=8,
                value=f'=COUNTIF(\'Review Queue\'!${status_letter}$5:${status_letter}${last_review_row},G{row_number})',
            )
            ws.cell(row=row_number, column=8).number_format = "#,##0"

    notes_start = max(17, 6 + len(metadata))
    ws.merge_cells(start_row=notes_start, start_column=1, end_row=notes_start, end_column=8)
    guidance = ws.cell(row=notes_start, column=1, value="Reviewer Guidance")
    guidance.fill = _fill(COLORS["teal_light"])
    guidance.font = Font(name=BODY_FONT, bold=True, color=COLORS["navy"])
    guidance.alignment = Alignment(vertical="center")
    for index, note in enumerate(notes, start=1):
        cell = ws.cell(row=notes_start + index, column=1, value=f"• {note}")
        cell.font = Font(name=BODY_FONT, size=10, color=COLORS["text"])

    files_start = notes_start + len(notes) + 3
    file_headers = [
        "Role",
        "Filename",
        "Load Policy",
        "Rows",
        "Columns",
        "Column Profiles",
        "Pair Profiles",
    ]
    for column, header in enumerate(file_headers, start=1):
        ws.cell(row=files_start, column=column, value=header)
    _style_header(ws, files_start, 1, len(file_headers))
    if files:
        for row_number, item in enumerate(files, start=files_start + 1):
            for column, header in enumerate(file_headers, start=1):
                cell = ws.cell(row=row_number, column=column, value=item.get(header))
                cell.font = Font(name=BODY_FONT, size=10, color=COLORS["text"])
                cell.alignment = Alignment(vertical="center")
                if column >= 4:
                    cell.number_format = "#,##0"
        files_end = files_start + len(files)
        _add_table(ws, f"A{files_start}:G{files_end}", "OverviewFilesTable")

    widths = {"A": 25, "B": 72, "C": 20, "D": 20, "E": 15, "F": 18, "G": 18, "H": 14}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A3"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def _validate_saved_workbook(path: Path, expected_sheets: list[str]) -> None:
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        raise WorkbookBuildError(f"Saved workbook cannot be reopened: {exc}") from exc
    try:
        if workbook.sheetnames != expected_sheets:
            raise WorkbookBuildError(
                f"Saved workbook sheet order mismatch: {workbook.sheetnames!r}"
            )
        if any(cell.data_type == "e" for ws in workbook.worksheets for row in ws.iter_rows() for cell in row):
            raise WorkbookBuildError("Saved workbook contains an Excel error cell")
    finally:
        workbook.close()


def _render_previews(workbook_path: Path, preview_dir: Path, sheet_names: list[str]) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    system_pdftoppm = Path("/usr/bin/pdftoppm")
    pdftoppm = str(system_pdftoppm) if system_pdftoppm.is_file() else shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise WorkbookBuildError(
            "--preview-dir requires LibreOffice (soffice) and pdftoppm on PATH"
        )
    preview_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".gwas-profile-preview-", dir=preview_dir
    ) as temp_name:
        temp = Path(temp_name)
        process_env = os.environ.copy()
        process_env["TMPDIR"] = str(temp)
        for index, sheet_name in enumerate(sheet_names):
            workbook = load_workbook(workbook_path)
            try:
                for ws in workbook.worksheets:
                    ws.sheet_state = "visible" if ws.title == sheet_name else "hidden"
                single_path = temp / f"sheet-{index + 1}.xlsx"
                workbook.save(single_path)
            finally:
                workbook.close()
            completed = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(single_path)],
                text=True,
                capture_output=True,
                check=False,
                env=process_env,
            )
            pdf_path = single_path.with_suffix(".pdf")
            if completed.returncode != 0 or not pdf_path.is_file():
                detail = (completed.stderr or completed.stdout).strip()
                raise WorkbookBuildError(f"Preview conversion failed for {sheet_name!r}: {detail}")
            slug = re.sub(r"[^a-z0-9]+", "-", sheet_name.lower()).strip("-")
            output_root = preview_dir / slug
            rendered = subprocess.run(
                [pdftoppm, "-f", "1", "-singlefile", "-png", "-r", "120", str(pdf_path), str(output_root)],
                text=True,
                capture_output=True,
                check=False,
                env=process_env,
            )
            if rendered.returncode != 0 or not output_root.with_suffix(".png").is_file():
                raise WorkbookBuildError(f"PNG rendering failed for {sheet_name!r}: {rendered.stderr.strip()}")


def build_workbook(
    payload: Mapping[str, Any],
    output_path: Path | str,
    preview_dir: Path | str | None = None,
) -> Path:
    """Build, validate, and atomically publish an XLSX workbook."""
    output = Path(output_path).resolve()
    if output.suffix.lower() != ".xlsx":
        raise WorkbookBuildError("output_path must end with .xlsx")
    if output.exists():
        raise WorkbookBuildError(f"Output already exists and will not be overwritten: {output}")
    specs = list(payload.get("sheets", []))
    names = [str(spec.get("name", "")) for spec in specs]
    if not names or any(not name for name in names) or len(names) != len(set(names)):
        raise WorkbookBuildError("payload.sheets must contain unique, non-empty names")

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    for name in names:
        workbook.create_sheet(name)
    workbook.properties.creator = "gwas-profile-report"
    workbook.properties.title = str(payload.get("title", "GWAS Profiling Review"))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    _populate_overview(overview, payload)
    for spec in specs:
        _populate_data_sheet(workbook[str(spec["name"])], spec)

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{output.stem}-", suffix=".part.xlsx", dir=output.parent, delete=False
    ) as handle:
        staged = Path(handle.name)
    try:
        workbook.save(staged)
        workbook.close()
        expected_sheets = ["Overview", *names]
        _validate_saved_workbook(staged, expected_sheets)
        if preview_dir is not None:
            _render_previews(staged, Path(preview_dir).resolve(), expected_sheets)
        os.replace(staged, output)
    except Exception:
        workbook.close()
        staged.unlink(missing_ok=True)
        raise

    return output


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a GWAS profiling XLSX from a JSON payload")
    parser.add_argument("payload", type=Path, help="Report payload JSON")
    parser.add_argument("output", type=Path, help="New XLSX output path")
    parser.add_argument("--preview-dir", type=Path, help="Optional PNG preview directory")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    with args.payload.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    build_workbook(payload, args.output, args.preview_dir)
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
